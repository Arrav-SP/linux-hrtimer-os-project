#!/usr/bin/env python3
"""
Build final_report.pdf and final_report.docx from report.md.

    python final_report/build_report.py

Requires: markdown-it-py, beautifulsoup4, python-docx, and Chrome or Edge
(used headless to print the HTML to PDF).
"""

import os
import re
import shutil
import subprocess
import sys

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "report.md")
HTML = os.path.join(HERE, "report.html")
PDF = os.path.join(HERE, "final_report.pdf")
DOCX = os.path.join(HERE, "final_report.docx")

CSS = """
@page { size: A4; margin: 2.0cm 2.1cm 2.1cm 2.1cm;
        @bottom-center { content: counter(page); font-family: Georgia, serif; font-size: 9pt; color: #444; } }
html { font-family: Georgia, 'Times New Roman', serif; font-size: 10.2pt; line-height: 1.4; color: #111; }
body { margin: 0; }
h1 { font-size: 19pt; line-height: 1.2; margin: 0 0 4pt 0; font-weight: bold; }
p.subtitle { font-size: 12.5pt; margin: 0 0 8pt 0; color: #222; }
p.meta { font-size: 9.2pt; color: #444; margin: 0 0 14pt 0; padding-bottom: 8pt; border-bottom: 0.6pt solid #888; text-align: left; }
h2 { font-size: 13pt; margin: 16pt 0 5pt 0; break-after: avoid; break-before: auto; }
p { break-before: auto; }
h3 { font-size: 11pt; margin: 11pt 0 3pt 0; break-after: avoid; }
p { margin: 0 0 6.5pt 0; text-align: justify; hyphens: auto; orphans: 2; widows: 2; }
ul, ol { margin: 0 0 6.5pt 0; padding-left: 18pt; }
li { margin-bottom: 2.5pt; text-align: justify; }
code { font-family: Consolas, 'Courier New', monospace; font-size: 8.8pt; }
pre { background: #f5f5f3; border: 0.5pt solid #ccc; padding: 5pt 7pt; margin: 4pt 0 8pt 0;
      font-size: 8.4pt; line-height: 1.3; white-space: pre-wrap; break-inside: avoid; }
pre code { font-size: 8.4pt; }
figure { margin: 8pt 0 9pt 0; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; }
figcaption, p.caption { font-size: 9pt; line-height: 1.32; text-align: left; margin-top: 3pt; color: #222; }
p.caption { margin: 3pt 0 10pt 0; break-before: avoid; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 0 0; font-size: 8.9pt; line-height: 1.28;
        font-family: 'Segoe UI', Arial, sans-serif; break-inside: avoid; }
th { border-top: 0.9pt solid #222; border-bottom: 0.6pt solid #222; padding: 2.5pt 4pt; text-align: left; font-weight: 600; }
td { padding: 2pt 4pt; vertical-align: top; white-space: normal; border-bottom: 0.25pt solid #ddd; }
tr:last-child td { border-bottom: 0.9pt solid #222; }
div.small table { font-size: 8pt; }
div.small td, div.small th { padding: 1.6pt 3pt; }
td:first-child, td:nth-child(2) { white-space: nowrap; }
table.wide td { white-space: normal !important; }
p.ref { font-size: 9.3pt; text-align: left; padding-left: 20pt; text-indent: -20pt; margin-bottom: 4pt; hyphens: none; overflow-wrap: anywhere; }
a { color: inherit; text-decoration: none; }
"""


def build_html():
    md = MarkdownIt("commonmark", {"html": True}).enable("table")
    body = md.render(open(SRC, encoding="utf-8").read())
    soup = BeautifulSoup(body, "html.parser")
    in_refs = False
    for el in soup.find_all(["h2", "p"]):
        if el.name == "h2":
            in_refs = el.get_text().strip() == "References"
            continue
        text = el.get_text()
        if re.match(r"^Table \d+\.", text):
            el["class"] = "caption"
        elif in_refs and re.match(r"^\[\d+\]", text):
            el["class"] = "ref"
    for tb in soup.find_all("table"):
        if len(tb.find("tr").find_all("th")) == 3:
            tb["class"] = "wide"
    html = ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
            "<title>Linux hrtimer case study</title><style>" + CSS + "</style></head><body>"
            + str(soup) + "</body></html>")
    open(HTML, "w", encoding="utf-8").write(html)
    return soup


def find_browser():
    for path in [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        shutil.which("google-chrome") or "",
        shutil.which("chromium") or "",
        shutil.which("chromium-browser") or "",
    ]:
        if path and os.path.exists(path):
            return path
    sys.exit("No Chrome/Edge/Chromium found for PDF rendering.")


def build_pdf():
    subprocess.run([
        find_browser(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={PDF}", "file:///" + HTML.replace("\\", "/"),
    ], check=True, capture_output=True, timeout=180)


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------

def add_inline(par, node, bold=False, italic=False, code=False, size=None):
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child).replace("\n", " ")
            if not text:
                continue
            run = par.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
            if code:
                run.font.name = "Consolas"
                run.font.size = Pt((size or 10.5) - 1.3)
            elif size:
                run.font.size = Pt(size)
        elif child.name in ("strong", "b"):
            add_inline(par, child, True, italic, code, size)
        elif child.name in ("em", "i"):
            add_inline(par, child, bold, True, code, size)
        elif child.name == "code":
            add_inline(par, child, bold, italic, True, size)
        elif child.name == "sup":
            run = par.add_run(child.get_text())
            run.font.superscript = True
        elif child.name == "br":
            par.add_run().add_break()
        else:
            add_inline(par, child, bold, italic, code, size)


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, sz in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "222222")
        borders.append(el)
    tcPr.append(borders)


def add_page_number(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


def build_docx(soup):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.3)
    add_page_number(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 19), ("Heading 1", 13), ("Heading 2", 11)):
        st = doc.styles[name]
        st.font.name = "Georgia"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
        st.element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        st.element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")

    text_width = sec.page_width - sec.left_margin - sec.right_margin

    def handle(el):
        if isinstance(el, NavigableString):
            return
        name = el.name
        cls = el.get("class") or []
        if isinstance(cls, str):
            cls = [cls]
        if name == "h1":
            doc.add_paragraph(el.get_text(), style="Title")
        elif name == "h2":
            doc.add_paragraph(el.get_text(), style="Heading 1")
        elif name == "h3":
            doc.add_paragraph(el.get_text(), style="Heading 2")
        elif name == "p":
            par = doc.add_paragraph()
            if "subtitle" in cls:
                add_inline(par, el, size=12.5)
            elif "meta" in cls:
                add_inline(par, el, size=9.2)
                par.paragraph_format.space_after = Pt(14)
            elif "caption" in cls:
                add_inline(par, el, size=9)
                par.paragraph_format.space_after = Pt(10)
            elif "ref" in cls:
                add_inline(par, el, size=9.3)
                par.paragraph_format.left_indent = Cm(0.7)
                par.paragraph_format.first_line_indent = Cm(-0.7)
                par.paragraph_format.space_after = Pt(4)
            else:
                add_inline(par, el)
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif name in ("ul", "ol"):
            style = "List Bullet" if name == "ul" else "List Number"
            for li in el.find_all("li", recursive=False):
                par = doc.add_paragraph(style=style)
                add_inline(par, li)
                par.paragraph_format.space_after = Pt(2.5)
        elif name == "pre":
            par = doc.add_paragraph()
            run = par.add_run(el.get_text().rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(8.4)
            par.paragraph_format.left_indent = Cm(0.4)
            par.paragraph_format.line_spacing = 1.0
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif name == "figure":
            img = el.find("img")
            width = text_width
            m = re.search(r"width:\s*(\d+)%", img.get("style", ""))
            if m:
                width = int(text_width * int(m.group(1)) / 100)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.paragraph_format.space_after = Pt(3)
            par.add_run().add_picture(os.path.join(HERE, img["src"]), width=width)
            cap = doc.add_paragraph()
            add_inline(cap, el.find("figcaption"), size=9)
            cap.paragraph_format.space_after = Pt(10)
        elif name == "table":
            rows = el.find_all("tr")
            ncol = len(rows[0].find_all(["th", "td"]))
            small = ncol > 8
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, tr in enumerate(rows):
                for c, cell_el in enumerate(tr.find_all(["th", "td"])):
                    cell = table.cell(r, c)
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
                    add_inline(par, cell_el, bold=(r == 0), size=7.5 if small else 8.6)
                    for run in par.runs:
                        if run.font.name != "Consolas":
                            run.font.name = "Segoe UI"
                    if r == 0:
                        set_cell_border(cell, top=8, bottom=6)
                    elif r == len(rows) - 1:
                        set_cell_border(cell, bottom=8)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif name == "div":
            for child in el.children:
                handle(child)

    for el in soup.children:
        handle(el)
    doc.save(DOCX)


if __name__ == "__main__":
    soup = build_html()
    build_pdf()
    build_docx(soup)
    print("Wrote", PDF)
    print("Wrote", DOCX)
